# accounts/views.py

from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse
from django.utils import timezone

from django.contrib.auth.forms import UserCreationForm
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin

from django.views.generic import ListView, CreateView, UpdateView, DeleteView, DetailView, View

from .models import StaffMember, ServiceHistory, DisciplinaryAction, AuditLog

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from io import BytesIO

from pathlib import Path
from datetime import datetime
import zipfile
import shutil
import tempfile

from django.conf import settings
from django.contrib import messages
from reportlab.lib.enums import TA_CENTER

# ---------------------------------------------------------------------------
# Auth: signup (staff/admin only) and home
# ---------------------------------------------------------------------------

@login_required
@user_passes_test(lambda u: u.is_staff)
def signup(request):
    if request.method == 'POST':
        form = UserCreationForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('staff_list')
    else:
        form = UserCreationForm()
    return render(request, 'signup.html', {'form': form})


@login_required
def home(request):
    return render(request, 'home.html')


# ---------------------------------------------------------------------------
# StaffMember: list, terminated list, detail
# ---------------------------------------------------------------------------

class StaffListView(LoginRequiredMixin, ListView):
    model = StaffMember
    template_name = 'staff_list.html'
    paginate_by = 10

    def get_queryset(self):
        query = self.request.GET.get('q')
        status = self.request.GET.get('status')
        gender = self.request.GET.get('gender')
        contract_type = self.request.GET.get('contract_type')
        posting_place = self.request.GET.get('posting_place')
        queryset = StaffMember.objects.all()

        if query:
            queryset = queryset.filter(name__icontains=query)

        if status:
            queryset = queryset.filter(status=status)

        if gender:
            queryset = queryset.filter(gender=gender)

        if contract_type:
            queryset = queryset.filter(contract_type=contract_type)
        if posting_place:
            queryset = queryset.filter(posting_place=posting_place)
        return queryset
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)

        context['posting_places'] = (
            StaffMember.objects
            .values_list('posting_place', flat=True)
            .distinct()
            .order_by('posting_place')
        )

        return context

class TerminatedStaffListView(LoginRequiredMixin, ListView):
    model = StaffMember
    template_name = 'terminated_staff_list.html'
    paginate_by = 10

    def get_queryset(self):
        return StaffMember.objects.exclude(status='active')


class StaffDetailView(LoginRequiredMixin, DetailView):
    model = StaffMember
    template_name = 'staff_detail.html'


# ---------------------------------------------------------------------------
# StaffMember: create, update, delete
# ---------------------------------------------------------------------------

STAFF_FIELDS = [
    'serial_number',
    'name',
    'designation',
    'pay_scale',
    'date_of_joining',
    'basic_pay',
    'posting_place',
    'gross_pay',
    'photo',
    'contract_type',
    'gender',
    'status',
    'separation_date',
]


'''class StaffCreateView(LoginRequiredMixin, CreateView):
    model = StaffMember
    fields = STAFF_FIELDS
    template_name = 'add_staff.html'
    success_url = '/accounts/staff/'

    def form_valid(self, form):
        response = super().form_valid(form)
        AuditLog.objects.create(
            user=self.request.user,
            action='created',
            model_name='StaffMember',
            object_repr=str(self.object),
        )
        return response '''
from django.contrib.auth.mixins import LoginRequiredMixin, PermissionRequiredMixin

class StaffCreateView(LoginRequiredMixin, PermissionRequiredMixin, CreateView):
    model = StaffMember
    fields = STAFF_FIELDS
    template_name = 'add_staff.html'
    success_url = '/accounts/staff/'
    permission_required = 'accounts.add_staffmember'

    def form_valid(self, form):
        response = super().form_valid(form)
        AuditLog.objects.create(
            user=self.request.user,
            action='created',
            model_name='StaffMember',
            object_repr=str(self.object),
        )
        return response

'''class StaffUpdateView(LoginRequiredMixin, UserPassesTestMixin, UpdateView):
    model = StaffMember
    fields = STAFF_FIELDS
    template_name = 'add_staff.html'
    success_url = '/accounts/staff/'

    def test_func(self):
        return self.request.user.is_staff

    def form_valid(self, form):
        response = super().form_valid(form)
        AuditLog.objects.create(
            user=self.request.user,
            action='updated',
            model_name='StaffMember',
            object_repr=str(self.object),
        )
        return response
'''
class StaffUpdateView(LoginRequiredMixin, PermissionRequiredMixin, UpdateView):
    model = StaffMember
    fields = STAFF_FIELDS
    template_name = 'add_staff.html'
    success_url = '/accounts/staff/'
    permission_required = 'accounts.change_staffmember'

    def form_valid(self, form):
        response = super().form_valid(form)
        AuditLog.objects.create(
            user=self.request.user,
            action='updated',
            model_name='StaffMember',
            object_repr=str(self.object),
        )
        return response
'''class StaffDeleteView(LoginRequiredMixin, UserPassesTestMixin, DeleteView):
    model = StaffMember
    template_name = 'confirm_delete_staff.html'
    success_url = '/accounts/staff/'

    def test_func(self):
        return self.request.user.is_superuser

    def delete(self, request, *args, **kwargs):
        self.object = self.get_object()
        AuditLog.objects.create(
            user=request.user,
            action='deleted',
            model_name='StaffMember',
            object_repr=str(self.object),
        )
        return super().delete(request, *args, **kwargs)
'''
class StaffDeleteView(LoginRequiredMixin, UserPassesTestMixin, DeleteView):
    model = StaffMember
    template_name = 'confirm_delete_staff.html'
    success_url = '/accounts/staff/'

    def test_func(self):
        return self.request.user.is_superuser

    def delete(self, request, *args, **kwargs):
        self.object = self.get_object()
        AuditLog.objects.create(
            user=request.user,
            action='deleted',
            model_name='StaffMember',
            object_repr=str(self.object),
        )
        return super().delete(request, *args, **kwargs)

# ---------------------------------------------------------------------------
# StaffMember: separate (soft-delete) and restore
# ---------------------------------------------------------------------------
'''class StaffSeparateView(LoginRequiredMixin, PermissionRequiredMixin, UpdateView):
   #class StaffSeparateView(LoginRequiredMixin, UserPassesTestMixin, UpdateView):
    model = StaffMember
    fields = ['status', 'separation_date']
    template_name = 'staff_separate.html'
    success_url = '/accounts/staff/'
    permission_required = 'accounts.change_staffmember'
    def test_func(self):
        return self.request.user.is_staff

    def form_valid(self, form):
        if not form.instance.separation_date:
            form.instance.separation_date = timezone.now().date()

        result = super().form_valid(form)
        ServiceHistory.objects.create(
            staff=self.object,
            position=f"Not in Service ({self.object.get_status_display()})",
            start_date=self.object.separation_date,
            end_date=None,
            salary=0,
            performance_rating='',
        )
        AuditLog.objects.create(
            user=self.request.user,
            action='separated',
            model_name='StaffMember',
            object_repr=str(self.object),
        )
        return result '''
class StaffSeparateView(LoginRequiredMixin, PermissionRequiredMixin, UpdateView):
    model = StaffMember
    fields = ['status', 'separation_date']
    template_name = 'staff_separate.html'
    success_url = '/accounts/staff/'
    permission_required = 'accounts.change_staffmember'

    def form_valid(self, form):
        if not form.instance.separation_date:
            form.instance.separation_date = timezone.now().date()

        result = super().form_valid(form)
        ServiceHistory.objects.create(
            staff=self.object,
            position=f"Not in Service ({self.object.get_status_display()})",
            start_date=self.object.separation_date,
            end_date=None,
            salary=0,
            performance_rating='',
        )
        AuditLog.objects.create(
            user=self.request.user,
            action='separated',
            model_name='StaffMember',
            object_repr=str(self.object),
        )
        return result

'''class StaffRestoreView(LoginRequiredMixin, UserPassesTestMixin, UpdateView):
    model = StaffMember
    fields = []
    template_name = 'staff_restore.html'
    success_url = '/accounts/staff/'

    def test_func(self):
        return self.request.user.is_staff

    def form_valid(self, form):
        self.object.status = 'active'
        self.object.save()

        gap_record = ServiceHistory.objects.filter(
            staff=self.object,
            position__icontains='Not in Service',
            end_date__isnull=True
        ).first()

        if gap_record:
            gap_record.end_date = timezone.now().date()
            gap_record.save()

        AuditLog.objects.create(
            user=self.request.user,
            action='restored',
            model_name='StaffMember',
            object_repr=str(self.object),
        )
        return redirect(self.success_url) '''

class StaffRestoreView(LoginRequiredMixin, PermissionRequiredMixin, UpdateView):
    model = StaffMember
    fields = []
    template_name = 'staff_restore.html'
    success_url = '/accounts/staff/'
    permission_required = 'accounts.change_staffmember'

    def form_valid(self, form):
        self.object.status = 'active'
        self.object.save()

        gap_record = ServiceHistory.objects.filter(
            staff=self.object,
            position__icontains='Not in Service',
            end_date__isnull=True
        ).first()

        if gap_record:
            gap_record.end_date = timezone.now().date()
            gap_record.save()

        AuditLog.objects.create(
            user=self.request.user,
            action='restored',
            model_name='StaffMember',
            object_repr=str(self.object),
        )
        return redirect(self.success_url)
# ---------------------------------------------------------------------------
# ServiceHistory: list, create, update, delete (scoped to one staff member)
# ---------------------------------------------------------------------------

class ServiceHistoryListView(LoginRequiredMixin, ListView):
    model = ServiceHistory
    template_name = 'staff_history.html'

    def get_queryset(self):
        self.staff = get_object_or_404(StaffMember, pk=self.kwargs['pk'])
        return ServiceHistory.objects.filter(staff=self.staff)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['staff'] = self.staff
        return context


class ServiceHistoryCreateView(LoginRequiredMixin, CreateView):
    model = ServiceHistory
    fields = ['position', 'start_date', 'end_date', 'salary', 'performance_rating']
    template_name = 'add_history.html'

    def form_valid(self, form):
        self.staff = get_object_or_404(StaffMember, pk=self.kwargs['pk'])
        form.instance.staff = self.staff
        return super().form_valid(form)

    def get_success_url(self):
        return f'/accounts/staff/{self.staff.pk}/history/'


class ServiceHistoryUpdateView(LoginRequiredMixin, UserPassesTestMixin, UpdateView):
    model = ServiceHistory
    fields = ['position', 'start_date', 'end_date', 'salary', 'performance_rating']
    template_name = 'add_history.html'

    def test_func(self):
        return self.request.user.is_staff

    def get_success_url(self):
        return f'/accounts/staff/{self.object.staff.pk}/history/'


class ServiceHistoryDeleteView(LoginRequiredMixin, UserPassesTestMixin, DeleteView):
    model = ServiceHistory
    template_name = 'confirm_delete_history.html'

    def test_func(self):
        return self.request.user.is_staff

    def get_success_url(self):
        return f'/accounts/staff/{self.object.staff.pk}/history/'


# ---------------------------------------------------------------------------
# DisciplinaryAction: list, create, update, delete (scoped to one staff member)
# ---------------------------------------------------------------------------

class DisciplinaryActionListView(LoginRequiredMixin, ListView):
    model = DisciplinaryAction
    template_name = 'disciplinary_list.html'

    def get_queryset(self):
        self.staff = get_object_or_404(StaffMember, pk=self.kwargs['pk'])
        return DisciplinaryAction.objects.filter(staff=self.staff)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['staff'] = self.staff
        return context


class DisciplinaryActionCreateView(LoginRequiredMixin, UserPassesTestMixin, CreateView):
    model = DisciplinaryAction
    fields = ['date', 'action_type', 'description']
    template_name = 'add_disciplinary.html'

    def test_func(self):
        return self.request.user.is_staff

    def form_valid(self, form):
        self.staff = get_object_or_404(StaffMember, pk=self.kwargs['pk'])
        form.instance.staff = self.staff
        response = super().form_valid(form)
        AuditLog.objects.create(
            user=self.request.user,
            action='created',
            model_name='DisciplinaryAction',
            object_repr=str(self.object),
        )
        return response

    def get_success_url(self):
        return f'/accounts/staff/{self.staff.pk}/disciplinary/'


class DisciplinaryActionUpdateView(LoginRequiredMixin, UserPassesTestMixin, UpdateView):
    model = DisciplinaryAction
    fields = ['date', 'action_type', 'description']
    template_name = 'add_disciplinary.html'

    def test_func(self):
        return self.request.user.is_staff

    def form_valid(self, form):
        response = super().form_valid(form)
        AuditLog.objects.create(
            user=self.request.user,
            action='updated',
            model_name='DisciplinaryAction',
            object_repr=str(self.object),
        )
        return response

    def get_success_url(self):
        return f'/accounts/staff/{self.object.staff.pk}/disciplinary/'


class DisciplinaryActionDeleteView(LoginRequiredMixin, UserPassesTestMixin, DeleteView):
    model = DisciplinaryAction
    template_name = 'confirm_delete_disciplinary.html'

    def test_func(self):
        return self.request.user.is_staff

    def delete(self, request, *args, **kwargs):
        self.object = self.get_object()
        AuditLog.objects.create(
            user=request.user,
            action='deleted',
            model_name='DisciplinaryAction',
            object_repr=str(self.object),
        )
        return super().delete(request, *args, **kwargs)

    def get_success_url(self):
        return f'/accounts/staff/{self.object.staff.pk}/disciplinary/'


# ---------------------------------------------------------------------------
# Reports: print view, Excel export, PDF export
# ---------------------------------------------------------------------------

class StaffListPrintView(LoginRequiredMixin, ListView):
    model = StaffMember
    template_name = 'staff_list_print.html'
    context_object_name = 'staff'

    def get_queryset(self):
        return StaffMember.objects.filter(status='active')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['report_date'] = timezone.now().strftime('%d %B %Y')
        return context

'''
class StaffExcelExportView(LoginRequiredMixin, View):
    def get(self, request):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Staff List"

        ws.append(['S#', 'Name', 'Designation', 'Pay Scale', 'Date of Joining', 'Posting Place', 'Status'])

        staff = StaffMember.objects.filter(status='active')
        for i, member in enumerate(staff, start=1):
            ws.append([
                i, member.name, member.designation, member.pay_scale,
                member.date_of_joining.strftime('%d-%m-%Y'),
                member.posting_place, member.get_status_display()
            ])

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="staff_list.xlsx"'
        wb.save(response)
        return response
'''
class StaffExcelExportView(LoginRequiredMixin, View):
    def get(self, request):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Staff List"

        ws.append([
            'S#',
            'Name',
            'Designation',
            'Pay Scale',
            'Date of Joining',
            'Posting Place',
            'Status'
        ])

        # Get filters from the URL
        query = request.GET.get('q')
        status = request.GET.get('status')
        gender = request.GET.get('gender')
        contract_type = request.GET.get('contract_type')

        # Start with all staff
        staff = StaffMember.objects.all()

        # Apply filters
        if query:
            staff = staff.filter(name__icontains=query)

        if status:
            staff = staff.filter(status=status)

        if gender:
            staff = staff.filter(gender=gender)

        if contract_type:
            staff = staff.filter(contract_type=contract_type)

        # Add filtered staff to Excel
        for i, member in enumerate(staff, start=1):
            ws.append([
                i,
                member.name,
                member.designation,
                member.pay_scale,
                member.date_of_joining.strftime('%d-%m-%Y'),
                member.posting_place,
                member.get_status_display()
            ])

        # Format header
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # Create borders
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Apply borders and alignment to all cells
        for row in ws.iter_rows():
            for cell in row:
                cell.border = thin_border
                cell.alignment = Alignment(
                    horizontal='center',
                    vertical='center',
                    wrap_text=True
                )

        # Adjust column widths
        column_widths = {
            'A': 8,
            'B': 25,
            'C': 25,
            'D': 12,
            'E': 18,
            'F': 25,
            'G': 15,
        }

        for column, width in column_widths.items():
            ws.column_dimensions[column].width = width

        # Adjust row heights
        ws.row_dimensions[1].height = 25

        for row in range(2, ws.max_row + 1):
            ws.row_dimensions[row].height = 22
        
        ws.freeze_panes = 'A2'
      

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

        response['Content-Disposition'] = 'attachment; filename="staff_list.xlsx"'

        wb.save(response)

        return response




class StaffPdfExportView(LoginRequiredMixin, View):
    def get(self, request):
        buffer = BytesIO()

        # Portrait page
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=20,
            leftMargin=20,
            topMargin=30,
            bottomMargin=30
        )

        styles = getSampleStyleSheet()

        elements = [
            Paragraph("Staff List Report", styles['Title'])
        ]

        # Normal table text
        table_style = ParagraphStyle(
            'TableStyle',
            parent=styles['Normal'],
            fontSize=8,
            leading=10,
            alignment=0
        )

        # Center aligned text
        center_style = ParagraphStyle(
            'CenterStyle',
            parent=styles['Normal'],
            fontSize=8,
            leading=10,
            alignment=TA_CENTER
        )

        # Header style
        header_style = ParagraphStyle(
            'HeaderStyle',
            parent=styles['Normal'],
            fontSize=8,
            leading=10,
            textColor=colors.white,
            alignment=TA_CENTER
        )

        data = [[
            Paragraph('S#', header_style),
            Paragraph('Name', header_style),
            Paragraph('Designation', header_style),
            Paragraph('Pay Scale', header_style),
            Paragraph('Date of Joining', header_style),
            Paragraph('Posting Place', header_style)
        ]]

        query = request.GET.get('q')
        status = request.GET.get('status')
        gender = request.GET.get('gender')
        contract_type = request.GET.get('contract_type')

        staff = StaffMember.objects.all()

        if query:
            staff = staff.filter(name__icontains=query)

        if status:
            staff = staff.filter(status=status)

        if gender:
            staff = staff.filter(gender=gender)

        if contract_type:
            staff = staff.filter(contract_type=contract_type)

        for i, member in enumerate(staff, start=1):

            data.append([
                Paragraph(str(i), table_style),

                Paragraph(
                    str(member.name or ''),
                    table_style
                ),

                Paragraph(
                    str(member.designation or ''),
                    table_style
                ),

                # Pay Scale - CENTER
                Paragraph(
                    str(member.pay_scale or ''),
                    center_style
                ),

                # Date of Joining - CENTER
                Paragraph(
                    member.date_of_joining.strftime('%d-%m-%Y')
                    if member.date_of_joining else '',
                    center_style
                ),

                Paragraph(
                    str(member.posting_place or ''),
                    table_style
                ),
            ])

        # Column widths
        col_widths = [
            25,     # S#
            105,    # Name
            105,    # Designation
            55,     # Pay Scale
            80,     # Date of Joining
            160     # Posting Place
        ]

        table = Table(
            data,
            colWidths=col_widths,
            repeatRows=1
        )

        table.setStyle(TableStyle([

            # Header
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a3c6e')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),

            # Grid
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),

            # Center Pay Scale + Date of Joining
            ('ALIGN', (3, 0), (4, -1), 'CENTER'),

            # Vertical alignment
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),

            # Padding
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))

        elements.append(table)

        doc.build(elements)

        buffer.seek(0)

        response = HttpResponse(
            buffer,
            content_type='application/pdf'
        )

        response['Content-Disposition'] = (
            'attachment; filename="staff_list.pdf"'
        )

        return response

from django.db.models import Count

class DashboardStatsView(LoginRequiredMixin, View):
    def get(self, request):
        active_staff = StaffMember.objects.filter(status='active')

        context = {
            'total_active': active_staff.count(),
            'total_terminated': StaffMember.objects.exclude(status='active').count(),
            'by_gender': active_staff.values('gender').annotate(count=Count('id')),
            'by_contract': active_staff.values('contract_type').annotate(count=Count('id')),
            'recent_separations': StaffMember.objects.exclude(status='active').order_by('-separation_date')[:5],
        }
        return render(request, 'dashboard_stats.html', context)
# ---------------------------------------------------------------------------
# Backup and Restore
# ---------------------------------------------------------------------------

class BackupView(LoginRequiredMixin, View):

    def get(self, request):

        if not request.user.is_superuser:
            return HttpResponse("Permission Denied", status=403)

        base_dir = Path(settings.BASE_DIR)

        database_file = base_dir / "db.sqlite3"
        media_folder = base_dir / "media"
        backup_folder = base_dir / "backups"

        if not database_file.exists():
            messages.error(request, "Database file not found.")
            return redirect('home')

        backup_folder.mkdir(exist_ok=True)

        timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")

        backup_filename = f"HR_Backup_{timestamp}.zip"
        backup_path = backup_folder / backup_filename

        with zipfile.ZipFile(
            backup_path,
            "w",
            zipfile.ZIP_DEFLATED
        ) as backup_zip:

            # Database
            backup_zip.write(
                database_file,
                arcname="db.sqlite3"
            )

            # Media folder
            if media_folder.exists():

                for file_path in media_folder.rglob("*"):

                    if file_path.is_file():

                        relative_path = file_path.relative_to(base_dir)

                        backup_zip.write(
                            file_path,
                            arcname=str(relative_path)
                        )

        messages.success(
            request,
            f"Backup created successfully: {backup_filename}"
        )

        return redirect('home')

class RestoreView(LoginRequiredMixin, View):

    def post(self, request):

        if not request.user.is_superuser:
            return HttpResponse("Permission Denied", status=403)

        backup_file = request.FILES.get('backup_file')

        if not backup_file:
            messages.error(request, "Please select a backup ZIP file.")
            return redirect('home')

        if not backup_file.name.lower().endswith('.zip'):
            messages.error(request, "Please select a ZIP backup file.")
            return redirect('home')

        base_dir = Path(settings.BASE_DIR)

        database_file = base_dir / "db.sqlite3"
        media_folder = base_dir / "media"
        backup_folder = base_dir / "backups"

        backup_folder.mkdir(exist_ok=True)

        # -------------------------------------------------
        # Save uploaded backup temporarily
        # -------------------------------------------------

        temp_backup = backup_folder / "restore_temp.zip"

        with open(temp_backup, "wb+") as destination:

            for chunk in backup_file.chunks():
                destination.write(chunk)

        # -------------------------------------------------
        # Validate ZIP
        # -------------------------------------------------

        if not zipfile.is_zipfile(temp_backup):

            temp_backup.unlink(missing_ok=True)

            messages.error(
                request,
                "Invalid backup file."
            )

            return redirect('home')

        # -------------------------------------------------
        # Validate database exists in ZIP
        # -------------------------------------------------

        with zipfile.ZipFile(temp_backup, "r") as backup_zip:

            file_list = backup_zip.namelist()

            if "db.sqlite3" not in file_list:

                temp_backup.unlink(missing_ok=True)

                messages.error(
                    request,
                    "Invalid backup: db.sqlite3 is missing."
                )

                return redirect('home')

            # Security check
            for file_name in file_list:

                target_path = (
                    base_dir / file_name
                ).resolve()

                if not str(target_path).startswith(
                    str(base_dir.resolve())
                ):

                    temp_backup.unlink(missing_ok=True)

                    messages.error(
                        request,
                        "Unsafe backup file detected."
                    )

                    return redirect('home')

        # -------------------------------------------------
        # Create emergency backup
        # -------------------------------------------------

        timestamp = datetime.now().strftime(
            "%Y-%m-%d_%H-%M-%S"
        )

        emergency_backup = (
            backup_folder /
            f"Emergency_Backup_{timestamp}.zip"
        )

        with zipfile.ZipFile(
            emergency_backup,
            "w",
            zipfile.ZIP_DEFLATED
        ) as emergency_zip:

            emergency_zip.write(
                database_file,
                arcname="db.sqlite3"
            )

            if media_folder.exists():

                for file_path in media_folder.rglob("*"):

                    if file_path.is_file():

                        relative_path = (
                            file_path.relative_to(base_dir)
                        )

                        emergency_zip.write(
                            file_path,
                            arcname=str(relative_path)
                        )

        # -------------------------------------------------
        # Extract and restore
        # -------------------------------------------------

        try:

            with tempfile.TemporaryDirectory() as temp_dir:

                temp_path = Path(temp_dir)

                with zipfile.ZipFile(
                    temp_backup,
                    "r"
                ) as backup_zip:

                    backup_zip.extractall(temp_path)

                restored_database = (
                    temp_path / "db.sqlite3"
                )

                if not restored_database.exists():

                    messages.error(
                        request,
                        "Restore failed: database missing."
                    )

                    return redirect('home')

                # Restore database
                shutil.copy2(
                    restored_database,
                    database_file
                )

                # Restore media
                restored_media = (
                    temp_path / "media"
                )

                if restored_media.exists():

                    if media_folder.exists():
                        shutil.rmtree(media_folder)

                    shutil.copytree(
                        restored_media,
                        media_folder
                    )

            messages.success(
                request,
                "System restored successfully."
            )

        except Exception as e:

            messages.error(
                request,
                f"Restore failed: {e}"
            )

        finally:

            temp_backup.unlink(missing_ok=True)

        return redirect('home')    
class BackupRestorePageView(LoginRequiredMixin, View):

    def get(self, request):

        if not request.user.is_superuser:
            return HttpResponse("Permission Denied", status=403)

        return render(
            request,
            'backup_restore.html'
        )    